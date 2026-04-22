"""
Military Medals Excel Builder — Wikipedia API
Fetches 40 medals (8 per country) with ribbon images from Wikipedia.
"""

import os
import re
import time
import requests
from io import BytesIO
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

# ── Config ────────────────────────────────────────────────────────────────────
API     = "https://en.wikipedia.org/w/api.php"
HEADERS = {"User-Agent": "MedalCatalogBot/1.0 (educational project)"}

IMAGE_DIR     = "ribbons"
EXCEL_FILE    = "medals_preview_40.xlsx"
LIMIT         = 8    # medals per country in output  → 5 × 8 = 40
FETCH         = 40   # candidates to scan per country
DELAY         = 1.5  # seconds between API calls

ROW_HEIGHT_PT = 22.5
COL_IMG_WIDTH = 15.0
IMG_W, IMG_H  = 110, 28

os.makedirs(IMAGE_DIR, exist_ok=True)

# ── Country sources ───────────────────────────────────────────────────────────
COUNTRIES = {
    "US": {"page": "Awards and decorations of the United States Armed Forces",
           "category": "Category:Military awards and decorations of the United States",
           "use_category": False},
    "CA": {"page": "Orders, decorations, and medals of Canada",
           "category": "Category:Military awards and decorations of Canada",
           "use_category": True},
    "NZ": {"page": "New Zealand honours order of wearing",
           "category": "Category:Military awards and decorations of New Zealand",
           "use_category": True},
    "AU": {"page": "Australian honours and awards system",
           "category": "Category:Military awards and decorations of Australia",
           "use_category": True},
    "IN": {"page": "Orders, decorations, and medals of India",
           "category": "Category:Military awards and decorations of India",
           "use_category": True},
}

TOP_MEDAL = {"US": "medal of honor", "CA": "victoria cross",
             "NZ": "victoria cross", "AU": "victoria cross", "IN": "param vir chakra"}

# ── Compiled regexes (module-level, compiled once) ────────────────────────────
RE_RIBBON_FNAME   = re.compile(r"ribbon|_bar\b|ribbon.bar|service.bar", re.I)
RE_SKIP_FNAME     = re.compile(
    r"flag|coat.of.arms|logo|icon|map|seal|portrait|wiki_letter|ambox|"
    r"question_mark|voting|emblem|insignia|crest|shield|star\.svg", re.I)
RE_JPEG           = re.compile(r"\.(jpg|jpeg)$", re.I)
RE_NEXT_ROW       = re.compile(r"next|higher|lower|preceded|followed|successor|predecessor", re.I)
RE_RIBBON_CAPTION = re.compile(r"\bribbon\b", re.I)
RE_INFOBOX        = re.compile(r"infobox", re.I)
RE_INFOBOX_LABEL  = re.compile(r"infobox-label", re.I)
RE_THUMB          = re.compile(r"/thumb/[0-9a-f]/[0-9a-f]{2}/([^/]+\.(svg|png|jpg|jpeg|gif))/", re.I)
RE_SAFE_NAME      = re.compile(r'[\\/*?:"<>|,\(\) ]')
RE_SKIP_MEDAL     = re.compile(
    r"^(List of|History of|Armed forces of|Military of|Comparison|"
    r"Wikipedia:|File:|Template:|Category:|\d{4}\s|Honours system|"
    r"honours system|awards system|order of wearing|order of precedence)", re.I)
RE_SKIP_CAT       = re.compile(
    r"^(List of|History of|Overview|System|Honours system)|"
    r"\bBadge\b|\bInsignia\b|\bUnit\b|\bSquadron\b|\bBattery\b|\bBattalion\b", re.I)
RE_SKIP_FALLBACK  = re.compile(r"\.(jpg|jpeg)$|flag|portrait|photo|emblem|logo|icon|seal|coat", re.I)

# ── API ───────────────────────────────────────────────────────────────────────

def api_get(params: dict) -> dict:
    params.setdefault("format", "json")
    params.setdefault("formatversion", "2")
    for attempt in range(3):
        try:
            r = requests.get(API, params=params, headers=HEADERS, timeout=20)
            if r.status_code == 429:
                time.sleep(6 * (attempt + 1))
                continue
            r.raise_for_status()
            return r.json()
        except Exception as e:
            print(f"  [API WARN] {e}")
            time.sleep(2)
    return {}

# ── Medal list fetchers ───────────────────────────────────────────────────────

def get_medal_links(page_title: str, limit: int) -> list[tuple[str, str]]:
    data  = api_get({"action": "query", "titles": page_title,
                     "prop": "links", "pllimit": "500", "plnamespace": "0"})
    pages = data.get("query", {}).get("pages", [])
    if not pages:
        return []
    results = []
    for lnk in pages[0].get("links", []):
        title = lnk.get("title", "")
        if RE_SKIP_MEDAL.match(title) or len(title) < 5:
            continue
        results.append((title, "https://en.wikipedia.org/wiki/" + title.replace(" ", "_")))
        if len(results) >= limit:
            break
    return results


def get_category_members(category: str, limit: int) -> list[tuple[str, str]]:
    data    = api_get({"action": "query", "list": "categorymembers",
                       "cmtitle": category, "cmtype": "page",
                       "cmlimit": str(limit * 6), "cmnamespace": "0"})
    members = data.get("query", {}).get("categorymembers", [])
    results = []
    for m in members:
        title = m.get("title", "")
        if RE_SKIP_CAT.search(title) or len(title) < 5:
            continue
        results.append((title, "https://en.wikipedia.org/wiki/" + title.replace(" ", "_")))
        if len(results) >= limit:
            break
    return results

# ── Ribbon finder ─────────────────────────────────────────────────────────────

def _real_filename(src: str) -> str:
    m = RE_THUMB.search(src)
    return requests.utils.unquote(m.group(1) if m else src.split("/")[-1])

def _is_bad(fname: str) -> bool:
    return bool(RE_SKIP_FNAME.search(fname) or RE_JPEG.search(fname))

def _ribbon_shape(img_tag) -> bool:
    try:
        fw, fh = int(img_tag.get("data-file-width", 0)), int(img_tag.get("data-file-height", 0))
        return fw > 0 and fh > 0 and fw / fh >= 2.5
    except (ValueError, TypeError):
        return False

def _resolve_file(fname: str, src: str) -> tuple[str, str]:
    """Call imageinfo API for a File: title, return (img_url, ribbon_page)."""
    if src.startswith("//"):
        img_url = "https:" + src
    elif src.startswith("http"):
        img_url = src
    else:
        img_url = "https://en.wikipedia.org" + src

    time.sleep(DELAY)
    data  = api_get({"action": "query", "titles": "File:" + fname,
                     "prop": "imageinfo", "iiprop": "url|mime|descriptionurl",
                     "iiurlwidth": "320"})
    pages = data.get("query", {}).get("pages", [])
    ribbon_page = ""
    if pages:
        info        = pages[0].get("imageinfo", [{}])[0]
        ribbon_page = info.get("descriptionurl", "")
        if not ribbon_page:
            canon       = pages[0].get("title", "File:" + fname)
            ribbon_page = "https://commons.wikimedia.org/wiki/" + canon.replace(" ", "_")
        thumb = info.get("thumburl", "")
        if thumb:
            img_url = thumb
    return img_url, ribbon_page


def get_ribbon(medal_title: str) -> tuple[str, str]:
    """
    Scan the Wikipedia page top-to-bottom, return the FIRST ribbon bar image.
    Accepts images that have a ribbon keyword in filename OR are ribbon-shaped,
    inside the infobox or a figure with a ribbon caption.
    Skips Next/Higher/Lower rows and JPEG photos.
    """
    data = api_get({"action": "parse", "page": medal_title, "prop": "text",
                    "disableeditsection": "true", "disabletoc": "true"})
    html = data.get("parse", {}).get("text", "")
    if not html:
        return _fallback(medal_title)

    soup    = BeautifulSoup(html, "html.parser")
    infobox = soup.find("table", class_=RE_INFOBOX)

    for img in soup.find_all("img"):
        src = img.get("src", "")
        if not src:
            continue
        fname = _real_filename(src)
        if _is_bad(fname):
            continue

        has_kw    = bool(RE_RIBBON_FNAME.search(fname))
        is_ribbon = has_kw or _ribbon_shape(img)
        if not is_ribbon:
            continue

        # Context check: infobox or ribbon-captioned figure
        in_infobox = bool(infobox and img.find_parent("table", class_=RE_INFOBOX))
        figure     = img.find_parent("figure")
        in_fig     = (figure is not None
                      and figure.find("figcaption") is not None
                      and RE_RIBBON_CAPTION.search(figure.find("figcaption").get_text(strip=True)))
        if not in_infobox and not in_fig:
            continue

        # Skip Next/Higher/Lower rows
        tr = img.find_parent("tr")
        if tr:
            lbl = tr.find(class_=RE_INFOBOX_LABEL)
            if lbl and RE_NEXT_ROW.search(lbl.get_text(strip=True)):
                continue

        # Skip primary medal photo (infobox, no ribbon keyword, not ribbon-shaped)
        if in_infobox and not has_kw and not _ribbon_shape(img):
            continue

        return _resolve_file(fname, src)

    return _fallback(medal_title)


def _fallback(medal_title: str) -> tuple[str, str]:
    """Last resort: scan prop=images for a ribbon-named file."""
    data  = api_get({"action": "query", "titles": medal_title,
                     "prop": "images", "imlimit": "50"})
    pages = data.get("query", {}).get("pages", [])
    if not pages:
        return "", ""

    title_words = set(re.findall(r"[a-z]{3,}", medal_title.lower()))
    best_file, best_score = None, 0
    for img in pages[0].get("images", []):
        ftitle = img.get("title", "")
        fname  = requests.utils.unquote(ftitle.replace("File:", ""))
        if RE_SKIP_FALLBACK.search(fname) or not RE_RIBBON_FNAME.search(fname):
            continue
        score = 3 + len(title_words & set(re.findall(r"[a-z]{3,}", fname.lower()))) * 2
        if score > best_score:
            best_score, best_file = score, ftitle

    if not best_file:
        return "", ""

    fname = requests.utils.unquote(best_file.replace("File:", ""))
    time.sleep(DELAY)
    data  = api_get({"action": "query", "titles": best_file,
                     "prop": "imageinfo", "iiprop": "url|mime|descriptionurl",
                     "iiurlwidth": "320"})
    pages = data.get("query", {}).get("pages", [])
    if not pages:
        return "", ""
    info     = pages[0].get("imageinfo", [{}])[0]
    thumb    = info.get("thumburl", "")
    original = info.get("url", "")
    mime     = info.get("mime", "")
    page_url = info.get("descriptionurl", "")
    if not page_url:
        canon    = pages[0].get("title", best_file)
        page_url = "https://commons.wikimedia.org/wiki/" + canon.replace(" ", "_")
    if thumb:
        return thumb, page_url
    if mime in ("image/png", "image/gif"):
        return original, page_url
    return "", page_url

# ── Image download ────────────────────────────────────────────────────────────

def download_image(url: str, filename: str) -> str | None:
    filename = RE_SAFE_NAME.sub("_", filename)
    path = os.path.join(IMAGE_DIR, filename)
    for attempt in range(3):
        try:
            r = requests.get(url, headers=HEADERS, timeout=20)
            if r.status_code == 429:
                wait = 6 * (attempt + 1)
                print(f"    [429] waiting {wait}s…")
                time.sleep(wait)
                continue
            if r.status_code != 200:
                print(f"    [WARN] HTTP {r.status_code}")
                return None
            pil = PILImage.open(BytesIO(r.content)).convert("RGBA")
            pil.thumbnail((IMG_W * 4, IMG_H * 4), PILImage.LANCZOS)
            pil.save(path, "PNG", optimize=True)
            return path
        except Exception as e:
            print(f"    [WARN] {e}")
            return None
    return None

# ── Collect medals ────────────────────────────────────────────────────────────
all_medals = []

for code, cfg in COUNTRIES.items():
    print(f"\n{'='*45}\n  {code} — {cfg['page']}\n{'='*45}")
    time.sleep(DELAY)

    medals = (get_category_members(cfg["category"], FETCH) if cfg["use_category"]
              else get_medal_links(cfg["page"], FETCH) or get_category_members(cfg["category"], FETCH))

    if not medals:
        print("  [WARN] no medals found")
        continue

    country_medals = []
    for idx, (name, wiki_url) in enumerate(medals, start=1):
        if len(country_medals) >= LIMIT:
            break
        print(f"  [{idx}] {name}")
        time.sleep(DELAY)

        img_url, ribbon_page = get_ribbon(name)
        if not img_url:
            print("       ✘ no ribbon — skipping")
            continue

        fname    = RE_SAFE_NAME.sub("_", f"{code}_{name}")[:80] + ".png"
        img_path = download_image(img_url, fname)
        print(f"       {'✔' if img_path else '✘ download failed'}  {img_url[:90]}")
        if not img_path:
            continue

        top      = TOP_MEDAL.get(code, "")
        priority = 1 if top in name.lower() else len(country_medals) + 1
        country_medals.append({"priority": priority, "country": code, "name": name,
                                "wiki_url": wiki_url, "ribbon_page": ribbon_page,
                                "img_path": img_path})

    # Fill remaining slots without images if needed
    if len(country_medals) < LIMIT:
        seen = {m["name"] for m in country_medals}
        for name, wiki_url in medals:
            if len(country_medals) >= LIMIT:
                break
            if name in seen:
                continue
            top      = TOP_MEDAL.get(code, "")
            priority = 1 if top in name.lower() else len(country_medals) + 1
            country_medals.append({"priority": priority, "country": code, "name": name,
                                   "wiki_url": wiki_url, "ribbon_page": "", "img_path": None})

    all_medals.extend(country_medals)

print(f"\nTotal medals collected: {len(all_medals)}")

# ── Build Excel ───────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Military Medals"

HDR_FILL   = PatternFill("solid", fgColor="1F3864")
HDR_FONT   = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
LINK_FONT  = Font(color="0563C1", underline="single", size=10, name="Calibri")
BODY_FONT  = Font(size=10, name="Calibri")
EVEN_FILL  = PatternFill("solid", fgColor="F0F4FF")
ODD_FILL   = PatternFill("solid", fgColor="FFFFFF")
CENTER     = Alignment(horizontal="center", vertical="center")
LEFT_MID   = Alignment(horizontal="left",   vertical="center")
HDR_BORDER = Border(bottom=Side(style="medium", color="AAAAAA"))

COLS = [("Priority", 8), ("Country", 9), ("Medal Name", 35),
        ("Wikipedia Link", 60), ("Ribbon Page Link", 60), ("Ribbon", COL_IMG_WIDTH)]

for ci, (label, width) in enumerate(COLS, 1):
    c = ws.cell(row=1, column=ci, value=label)
    c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CENTER, HDR_BORDER
    ws.column_dimensions[get_column_letter(ci)].width = width
ws.row_dimensions[1].height = 24

for ri, m in enumerate(all_medals, start=2):
    ws.row_dimensions[ri].height = ROW_HEIGHT_PT
    fill = EVEN_FILL if ri % 2 == 0 else ODD_FILL

    def cell(col, value, font=None, hyperlink=None, align=LEFT_MID):
        c = ws.cell(ri, col, value)
        c.fill, c.font, c.alignment = fill, font or BODY_FONT, align
        if hyperlink:
            c.hyperlink = hyperlink
        return c

    cell(1, m["priority"], align=CENTER)
    cell(2, m["country"],  align=CENTER)
    cell(3, m["name"])
    cell(4, m["wiki_url"],    font=LINK_FONT, hyperlink=m["wiki_url"])
    cell(5, m["ribbon_page"] or "N/A",
         font=LINK_FONT if m["ribbon_page"] else BODY_FONT,
         hyperlink=m["ribbon_page"] if m["ribbon_page"] else None)

    if m["img_path"] and os.path.exists(m["img_path"]):
        try:
            img = ExcelImage(m["img_path"])
            img.width, img.height = IMG_W, IMG_H
            ws.add_image(img, f"F{ri}")
        except Exception as e:
            print(f"  [WARN] embed failed: {e}")
            cell(6, "[embed error]")
    else:
        cell(6, "[no image]")

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"
wb.save(EXCEL_FILE)
print(f"\n✔  Excel  →  {EXCEL_FILE}")
print(f"✔  Ribbons →  {IMAGE_DIR}/")
